import openpyxl, os, sys

def cell_text(cell):
    """返回单元格文本，处理数字格式、换行与 None。"""
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, str):
        return v.replace("\r", " ").replace("\n", " ").strip()
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        nf = (cell.number_format or "")
        if ".00" in nf:
            return f"{v:.2f}"
        if v.is_integer():
            return str(int(v))
        return str(v)
    return str(v).strip()

def md_escape(s):
    # 转义管道符与 HTML 尖括号，避免 Markdown 表格错位及 <RED> 等标签被吞
    return s.replace("|", "\\|").replace("<", "&lt;").replace(">", "&gt;")

def convert(xlsx_path, out_dir=None):
    xlsx_path = os.path.abspath(xlsx_path)
    if not os.path.isfile(xlsx_path):
        print("SKIP (not found):", xlsx_path)
        return None
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    lines = []

    # 标题行 (row 1, 可能合并 A1:F1，可能含 \n 说明)
    title_raw = ws.cell(1, 1).value or ""
    title_parts = str(title_raw).split("\n")
    title_main = title_parts[0].strip()
    title_desc = "\n".join(p.strip() for p in title_parts[1:] if p.strip())
    lines.append(f"# {title_main}")
    lines.append("")
    if title_desc:
        lines.append(title_desc)
        lines.append("")

    # 表头 (row 2)
    ncol = ws.max_column
    headers = [cell_text(ws.cell(2, c)) for c in range(1, ncol + 1)]

    def emit_table(rows):
        lines.append("| " + " | ".join(md_escape(h) for h in headers) + " |")
        lines.append("| " + " | ".join("---" for _ in headers) + " |")
        for row in rows:
            lines.append("| " + " | ".join(md_escape(x) for x in row) + " |")
        lines.append("")

    cur_rows = []
    for r in range(3, ws.max_row + 1):
        a = cell_text(ws.cell(r, 1))
        b = cell_text(ws.cell(r, 2))
        if not a and not b:
            continue
        if a and not b:
            # 分类小标题行（A 有值、B 空）→ ## ；先输出上一段表格
            if cur_rows:
                emit_table(cur_rows)
                cur_rows = []
            lines.append(f"## {a}")
            lines.append("")
        else:
            row = [cell_text(ws.cell(r, c)) for c in range(1, ncol + 1)]
            cur_rows.append(row)

    if cur_rows:
        emit_table(cur_rows)

    out = "\n".join(lines).rstrip() + "\n"
    out_dir = out_dir or os.path.dirname(xlsx_path)
    os.makedirs(out_dir, exist_ok=True)
    out_md = os.path.join(out_dir, os.path.splitext(os.path.basename(xlsx_path))[0] + ".md")
    with open(out_md, "w", encoding="utf-8") as f:
        f.write(out)
    print("WROTE:", out_md, "chars=", len(out))
    return out_md

def main():
    args = sys.argv[1:]
    if not args:
        print("用法: python xlsx_to_md.py 文件1.xlsx [文件2.xlsx ...] [--out 输出目录]")
        print("说明: 将 BOQ/清单类 xlsx 转为 Markdown（标题→#；分类小标题→##；明细→表格）。")
        sys.exit(1)
    out_dir = None
    files = []
    i = 0
    while i < len(args):
        a = args[i]
        if a == "--out":
            out_dir = args[i + 1]
            i += 2
            continue
        elif a.startswith("--out="):
            out_dir = a.split("=", 1)[1]
            i += 1
            continue
        else:
            files.append(a)
            i += 1
    if not files:
        print("错误: 未提供任何 xlsx 文件")
        sys.exit(1)
    for f in files:
        convert(f, out_dir)

if __name__ == "__main__":
    main()
